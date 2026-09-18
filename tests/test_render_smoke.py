from pathlib import Path

from pptx import Presentation

from argenerator.layout import compute_layout
from argenerator.render_pptx import render_roadmap
from argenerator.schema import load_roadmap
from argenerator.theme import Theme

EXAMPLE_ROADMAP = Path(__file__).parent.parent / "examples" / "fictional_company_roadmap.yaml"


def _all_text(slide) -> str:
    chunks = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            chunks.append(shape.text_frame.text)
    return "\n".join(chunks)


def test_render_example_roadmap_end_to_end(tmp_path):
    roadmap = load_roadmap(EXAMPLE_ROADMAP)
    layout = compute_layout(roadmap)
    output_path = tmp_path / "roadmap.pptx"

    render_roadmap(layout, output_path, Theme(), EXAMPLE_ROADMAP)

    assert output_path.exists()

    presentation = Presentation(str(output_path))
    assert len(presentation.slides) == 1

    slide = presentation.slides[0]
    assert len(slide.shapes) > 20  # background, header cells, badges, tracks, markers, labels

    text = _all_text(slide)
    assert "Architecture Roadmap" in text
    assert "Customer Platform" in text
    assert "Kick off Storefront Rebuild" in text
    assert "Virtual Assistant Pilot Live" in text
    assert "Highly Confidential" in text
